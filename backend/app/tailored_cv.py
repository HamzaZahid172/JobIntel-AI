"""Build a reviewable, evidence-only CV draft and a portable Word document."""

import re
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt


_SECTION = re.compile(
    r"(?im)^\s*(?:professional summary|profile|summary|professional experience|"
    r"experience|technical skills|skills|education|projects|key projects)\s*$"
)


def build_tailored_cv_text(cv_text: str, draft: dict) -> str:
    """Insert grounded highlights while keeping the original CV text unchanged.

    Skills and fragments come from cv_tailoring_draft, which checks source
    evidence. Personal details, dates and all original sections remain present.
    The user can edit the resulting text before exporting it.
    """
    source = (cv_text or "").strip()
    if not source:
        return ""
    normalized = re.sub(r"\s+", " ", source).casefold()
    skills = [s for s in draft.get("matched_skills", [])
              if re.search(r"(?<!\w)" + re.escape(s) + r"(?!\w)", source, re.I)]
    evidence = [e for e in draft.get("cv_evidence", [])
                if e and re.sub(r"\s+", " ", e).casefold() in normalized]
    if not skills and not evidence:
        return source

    blocks = ["ROLE-RELEVANT HIGHLIGHTS"]
    if skills:
        # A mention can be negated or refer to a planned skill. Do not turn
        # keyword occurrence into a claim of proficiency.
        blocks.append("Job-related terms mentioned in CV: " + ", ".join(skills))
    blocks.extend("• " + item for item in evidence[:3])
    highlights = "\n".join(blocks)
    section = _SECTION.search(source)
    if section:
        pos = section.start()
    else:
        pos = source.find("\n")
        if pos < 0:
            pos = len(source)
    # Keep both slices byte-for-byte so the original remains recoverable.
    return source[:pos] + "\n\n" + highlights + "\n\n" + source[pos:]


def build_tailored_cv_docx(text: str) -> bytes:
    """Export edited text; this does not save, attach or submit the CV."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    for index, line in enumerate(text.splitlines()):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line)
        if index == 0 and len(line) <= 100:
            run.bold = True
            run.font.size = Pt(15)
        elif line and len(line) <= 70 and line.upper() == line and re.search(r"[A-Z]", line):
            run.bold = True
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
