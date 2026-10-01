import os
import re
from datetime import date
from io import BytesIO

import httpx
from docx import Document
from docx.shared import Pt

from .intelligence import profile_role_families, rank_cv_skills


async def generate_cover_letter_text(cv_text: str, job, match: dict, user) -> tuple[str, str]:
    matched = match.get("matched_skills") or []
    missing = match.get("missing_skills") or []
    top_cv = [row["skill"] for row in rank_cv_skills(cv_text, limit=8)]
    mode = "template"

    if os.getenv("USE_OLLAMA", "false").lower() == "true":
        try:
            prompt = f"""Write a concise, professional English cover letter for the candidate below.

Rules:
- 280 to 360 words.
- Use only facts supported by the CV text.
- Do not invent employers, achievements, certifications, languages, years, or tools.
- Focus on the strongest overlap with the job.
- Do not claim missing skills as existing skills.
- Mention the exact company and role.
- No placeholders.
- End with a professional closing using the candidate's name.

Candidate name: {user.display_name}
Candidate email: {user.email}
Role: {job.title}
Company: {job.company}
Location: {job.location}
Matched skills: {", ".join(matched)}
Missing skills: {", ".join(missing)}

CV:
{cv_text[:12000]}

Job description:
{job.description[:12000]}
"""
            async with httpx.AsyncClient(timeout=50) as client:
                response = await client.post(
                    f"{os.getenv('OLLAMA_URL', 'http://localhost:11434')}/api/generate",
                    json={
                        "model": os.getenv("OLLAMA_MODEL", "llama3.2:3b"),
                        "prompt": prompt,
                        "stream": False,
                        "options": {"temperature": 0.25},
                    },
                )
                response.raise_for_status()
                text = response.json().get("response", "").strip()
                if len(text.split()) >= 120:
                    return text, "ollama"
        except Exception:
            pass

    skill_phrase = ", ".join(matched[:6] or top_cv[:6])
    role_families = ", ".join(x.replace("_", " ") for x in profile_role_families(cv_text)[:3])
    opening = (
        f"Dear Hiring Team at {job.company},\n\n"
        f"I am writing to apply for the {job.title} position. "
        f"My background aligns with several of the role's core technical requirements, particularly {skill_phrase}."
    )
    body = (
        f"\n\nMy CV reflects hands-on experience across {role_families or 'software engineering and technical delivery'}. "
        "I have worked on building, testing, integrating and improving software systems, with a focus on reliable implementation "
        "and practical problem solving. The overlap between my experience and this role makes the opportunity especially relevant to me."
    )
    if matched:
        body += (
            "\n\nFor this position, the strongest match comes from "
            + ", ".join(matched[:5])
            + ". I would bring that existing experience to the role while continuing to deepen the areas that are specific to your environment."
        )
    closing = (
        f"\n\nI would welcome the opportunity to discuss how my experience can contribute to {job.company}. "
        "Thank you for considering my application.\n\nKind regards,\n"
        f"{user.display_name}\n{user.email}"
    )
    return opening + body + closing, mode


def build_cover_letter_docx(text: str, user, job) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    doc.add_paragraph(user.display_name)
    doc.add_paragraph(user.email)
    doc.add_paragraph(date.today().strftime("%d %B %Y"))
    doc.add_paragraph()
    heading = doc.add_paragraph()
    run = heading.add_run(f"Application for {job.title} at {job.company}")
    run.bold = True

    for block in re.split(r"\n\s*\n", text.strip()):
        paragraph = doc.add_paragraph(block.strip())
        paragraph.paragraph_format.space_after = Pt(8)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def safe_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_")[:80] or "cover_letter"
